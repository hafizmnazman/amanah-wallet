"""Plan task E5: builds the 11-slide deck, with a talk track in each slide's speaker notes.

    python scripts/build_deck.py                      # docs/deck/amanah-deck.pptx
    python scripts/build_deck.py --pdf                # and docs/deck/amanah-deck.pdf, through LibreOffice
    python scripts/build_deck.py --team "Team name" --members "Name One, Name Two"

Each slide makes one claim and shows one exhibit, so it reads without a presenter. Numbers come from the
files the runs write: redteam/results/final/summary.json, evidence/demo-summary.json,
evidence/transactions.csv and contracts/deployments/84532.json. A number whose file does not exist yet is
printed as a [[fill: ...]] marker and listed at the end, so rerun this after the runs. Needs python-pptx
(scripts/requirements.txt).
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import results  # noqa: E402

REPO = results.REPO
OUT = REPO / "docs" / "deck" / "amanah-deck.pptx"

# The suites' counts on the build this deck describes; scripts/test_all.sh prints them.
TESTS = {"contracts": 112, "backend": 141, "red team": 29, "release scripts": 32, "frontend unit": 25,
         "browser checks": 17}
# What run_demo.py records when a live planner refused before any signer, so nothing was sent (decision E-21)
PLANNER_STOPS = ("the planner refused", "the planner made no payout")

INK = RGBColor(0x1D, 0x23, 0x21)
MUTED = RGBColor(0x5B, 0x66, 0x61)
GREEN = RGBColor(0x0B, 0x5D, 0x3B)
PALE = RGBColor(0xEC, 0xF3, 0xEF)
LINE = RGBColor(0xC9, 0xD6, 0xCF)
RED = RGBColor(0xA8, 0x23, 0x1C)
PALE_RED = RGBColor(0xF8, 0xE9, 0xE7)
AMBER = RGBColor(0x8A, 0x5A, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)


# Facts from the runs


def fill(hint: str) -> str:
    return f"[[fill: {hint}]]"


def big(value: str) -> int:
    """The headline numbers' size: smaller while a value is still a fill marker."""
    return 24 if value.startswith("[[fill") else 54


def facts(team: str | None, members: str | None) -> dict[str, str]:
    f: dict[str, str] = {"team": team or fill("team name"), "members": members or fill("member names"),
                         "tests_total": str(sum(TESTS.values())),
                         "tests_split": ", ".join(f"{name} {n}" for name, n in TESTS.items())}

    rt = results.redteam_summary() or {}
    success, refusal = rt.get("attack_success", {}), rt.get("false_refusal", {})
    for arm in ("S1", "S2", "S3", "S4", "S5"):
        f[arm] = results.rate(success.get(arm)) or fill(f"{arm} attack success")
        f[f"{arm}_short"] = results.rate(success.get(arm), ci=False) or fill(f"{arm} %")
        f[f"FR{arm[1]}"] = results.rate(refusal.get(arm), ci=False) or fill(f"{arm} false refusals")
    # Payment attacks only: an extraction leak is a success in every arm but signs nothing (decision C-12)
    f["S1_bad"] = results.rate(rt.get("attack_success_transfer", {}).get("S1"), ci=False) or fill("S1 % of payment attacks")
    moved = rt.get("moved_funds")
    f["S5_moved"] = str(moved["k"]) if moved and moved.get("n") else fill("attacks that moved funds")
    joint = rt.get("joint_failure_planner_auditor")
    joint_ci = rt.get("joint_failure_planner_auditor_ci")
    f["joint"] = (f"{100 * joint:.1f}% (95% CI {100 * joint_ci[0]:.1f} to {100 * joint_ci[1]:.1f})"
                  if joint is not None and joint_ci else fill("joint failure, planner and Qwen auditor"))
    f["n_attacks"] = str(rt["n_attacks"]) if rt else fill("attack cases")
    f["n_legit"] = str(rt["n_legitimate"]) if rt else fill("legitimate cases")
    f["protocol_commit"] = results.protocol_commit() or fill("protocol commit")

    record = results.deployment()
    f["vault"] = results.short(record["vault"]) if record and record.get("vault") else fill("vault address")
    on_base = results.transactions()
    f["tx_count"] = str(len(on_base)) if on_base else fill("Base Sepolia transactions")

    outcomes = results.outcomes()
    stopped = []
    for n in range(1, 17):
        o = outcomes.get(str(n))
        f[f"ok{n}"] = ("passed" if o["ok"] else "FAILED") if o else fill("run")
        if o and o["tx_hashes"]:
            f[f"tx{n}"] = results.short(o["tx_hashes"][-1])
        elif o and o["ok"] and any(p in o.get("got", "") for p in PLANNER_STOPS):
            f[f"tx{n}"] = "no transaction"
            stopped.append(str(n))
        else:
            f[f"tx{n}"] = fill(f"scenario {n} tx")
    f["planner_stopped"] = " and ".join(stopped)
    e1 = outcomes.get("E1")
    found = re.search(r"(\d+) of (\d+) pass", e1["got"]) if e1 else None
    f["verified"] = f"{found.group(1)} of {found.group(2)}" if found else fill("payouts verify.py passed")
    f["scenarios_passed"] = (str(sum(outcomes[str(n)]["ok"] for n in range(1, 17) if str(n) in outcomes))
                             if outcomes else fill("scenarios passed"))
    return f


# Drawing


def text_box(slide, left, top, width, height, paragraphs, *, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=1.1, after=Pt(7)):
    """`paragraphs` is a string or a list of strings; a list item that is a tuple is (text, size, color, bold)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    items = [paragraphs] if isinstance(paragraphs, str) else paragraphs
    for i, item in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        words, s, c, b = item if isinstance(item, tuple) else (item, size, color, bold)
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = after if i < len(items) - 1 else Pt(0)
        run = p.add_run()
        run.text = words
        run.font.name, run.font.size, run.font.color.rgb, run.font.bold = FONT, Pt(s), c, b
    return box


def panel(slide, left, top, width, height, *, fill_color=PALE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    s.text_frame.text = ""
    return s


def boxed(slide, left, top, width, height, lines, *, fill_color=PALE, line=None, size=14, color=INK,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    panel(slide, left, top, width, height, fill_color=fill_color, line=line)
    pad = Inches(0.18)
    return text_box(slide, left + pad, top + Inches(0.1), width - 2 * pad, height - Inches(0.2), lines, size=size,
                    color=color, align=align, anchor=anchor)


def arrow(slide, left, top, width=Inches(0.32), height=Inches(0.3), color=GREEN):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def table(slide, left, top, width, rows, widths, *, size=13, header_size=13, row_height=Inches(0.5)):
    """rows[0] is the header. widths are fractions of `width`."""
    shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, row_height * len(rows))
    t = shape.table
    for i, frac in enumerate(widths):
        t.columns[i].width = Emu(int(width * frac))
    for r, row in enumerate(rows):
        t.rows[r].height = row_height
        for c, value in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = GREEN if r == 0 else (WHITE if r % 2 else PALE)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame = cell.text_frame
            frame.word_wrap = True
            frame.text = ""
            run = frame.paragraphs[0].add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(header_size if r == 0 else size)
            run.font.bold = r == 0 or c == 0
            run.font.color.rgb = WHITE if r == 0 else INK
    return shape


class Deck:
    def __init__(self, f: dict[str, str]):
        self.f = f
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.count = 0

    def s(self, text: str) -> str:
        return text.format_map(self.f)

    def slide(self, kicker: str, claim: str, source: str, notes: str):
        self.count += 1
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        text_box(slide, MARGIN, Inches(0.38), W - 2 * MARGIN, Inches(0.35), self.s(kicker).upper(), size=12,
                 color=GREEN, bold=True)
        text_box(slide, MARGIN, Inches(0.72), W - 2 * MARGIN, Inches(1.2), self.s(claim), size=28, bold=True,
                 spacing=1.0)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, H - Inches(0.62), W - 2 * MARGIN, Pt(0.75))
        line.fill.solid()
        line.fill.fore_color.rgb = LINE
        line.line.fill.background()
        text_box(slide, MARGIN, H - Inches(0.55), W - 2 * MARGIN - Inches(1.6), Inches(0.4), "Source: " + self.s(source),
                 size=10, color=MUTED)
        text_box(slide, W - MARGIN - Inches(1.4), H - Inches(0.55), Inches(1.4), Inches(0.4),
                 f"Amanah  {self.count} / 11", size=10, color=MUTED, align=PP_ALIGN.RIGHT)
        slide.notes_slide.notes_text_frame.text = self.s(notes)
        return slide

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)


TOP = Inches(2.05)  # where an exhibit starts, below a two-line claim


def build(f: dict[str, str]) -> Deck:
    d = Deck(f)

    # 1. Title and the headline result
    s = d.slide(
        "Agent wallet · money held in trust",
        "Amanah: the agent wallet with no single point of persuasion",
        "red-team final run, arms S1 and S5 (redteam/results/final/report.md)",
        "Amanah is an agent wallet for money held in trust: zakat, charity, humanitarian cash and public "
        "subsidies. It pays beneficiaries on its own, on a public chain. Our claim is that no single persuaded "
        "model can make it pay. In our red team, {S1_bad} of the attacks that asked for a payment got the planner "
        "alone, which is how a typical agent wallet works, to sign it. Through Amanah's quorum and contract, {S5_moved} "
        "moved any funds. Amanah is the Arabic word for a trust held on behalf of others. The rest of this deck "
        "shows how the wallet carries that duty, and how you can check each number yourself.")
    text_box(s, MARGIN, Inches(1.75), Inches(11.5), Inches(0.6),
             "Three signers built to fail differently must agree, and the contract checks that they differ.",
             size=18, color=MUTED)
    boxed(s, MARGIN, Inches(2.8), Inches(5.85), Inches(2.6),
          [(d.s("{S1_bad}"), big(d.s("{S1_bad}")), RED, True),
           ("of red-team payment attacks got the planner alone, a typical single-model agent wallet, to sign a "
            "bad payment", 16, INK, False)], fill_color=PALE_RED)
    boxed(s, MARGIN + Inches(6.25), Inches(2.8), Inches(5.85), Inches(2.6),
          [(d.s("{S5_moved}"), big(d.s("{S5_moved}")), GREEN, True),
           ("payment attacks moved any funds through Amanah's quorum and contract, on the same set", 16, INK, False)])
    text_box(s, MARGIN, Inches(5.75), Inches(12), Inches(0.8),
             [(d.s("Team {team} · Universiti Teknologi PETRONAS, Malaysia"), 16, INK, True),
              (d.s("{members}"), 14, MUTED, False)])

    # 2. The problem
    s = d.slide(
        "1 · The problem",
        "Prompt injection is persuasion, and most agent wallets have one model to persuade",
        "README section 1: The Block, FBI and NCC Group, Giskard and OECD, Zscaler, ICML 2025",
        "Agents already hold money, and they are already being talked into moving it. Freysa released its prize "
        "pool after one message redefined its transfer tool. Bybit's signers approved a malicious payload while "
        "the screen showed the right address. An encoded message made an AI wallet send tokens worth about "
        "150,000 to 200,000 US dollars. Production wallets add allowlists, caps and session limits. Those are "
        "necessary, but they judge mechanics: who, how much and how often. A payment to a known address, under "
        "the limit, made for the wrong reason, passes all of them. And a model cannot reliably check another of "
        "its kind: when two models are both wrong, they pick the same wrong answer about 60% of the time, more "
        "often when they share a provider.")
    table(s, MARGIN, TOP, Inches(8.1), [
        ["When", "What happened", "Loss"],
        ["Nov 2024", "Freysa released its prize pool when one message redefined its transfer tool", "about US$47K"],
        ["Feb 2025", "Bybit's signers approved a malicious payload while the screen showed the right address",
         "about US$1.5B"],
        ["May 2026", "An NFT gave an AI wallet rights; an encoded message then made it send 3B DRB",
         "about US$150K to 200K"],
        ["Jul 2026", "Live websites hid payment instructions for agents; 4 of 26 models paid", "sandbox test"],
    ], [0.15, 0.62, 0.23], size=14, row_height=Inches(0.85))
    boxed(s, MARGIN + Inches(8.45), TOP, Inches(3.65), Inches(4.25),
          [("Why existing controls miss it", 16, GREEN, True),
           ("Policy engines check who, how much and how often.", 15, INK, False),
           ("An injected payout to a known address, under the cap, passes all three.", 15, INK, False),
           ("When two models are both wrong, they pick the same wrong answer about 60% of the time (ICML 2025).",
            15, INK, False)], anchor=MSO_ANCHOR.TOP)

    # 3. Money held in trust
    s = d.slide(
        "2 · Who it is for",
        "For money held in trust, the payer, the owner and the recipient are three different people",
        "README section 2: Utusan, ALNAP, WFP",
        "Every agent payment rail we reviewed serves the person who owns the money. None we reviewed is built for "
        "money held in trust. There the payer is a steward, such as a zakat institution or an aid agency. The owner is the "
        "public that gave the money. And the recipient is vulnerable and cannot absorb a mistake. A misdirected "
        "payment is money taken from someone poor, and it costs the institution the trust it runs on. The scale is "
        "large: one Malaysian state collected 1.22 billion ringgit of zakat in 2024, and humanitarian cash reached "
        "7.3 billion US dollars in 2025. Aid is already moving on-chain.")
    y = TOP
    for i, (title, body) in enumerate([("The payer", "a steward: a zakat institution, a charity, an aid agency"),
                                       ("The owner", "the public that gave the money"),
                                       ("The recipient", "vulnerable, and cannot absorb a mistake")]):
        x = MARGIN + i * Inches(4.15)
        boxed(s, x, y, Inches(3.75), Inches(1.45), [(title, 18, GREEN, True), (body, 15, INK, False)])
        if i < 2:
            arrow(s, x + Inches(3.8), y + Inches(0.58))
    table(s, MARGIN, y + Inches(1.85), W - 2 * MARGIN, [
        ["Fiduciary money", "Scale"],
        ["Zakat in Selangor, Malaysia, one state", "RM1.22B collected in 2024, reaching more than 370,000 asnaf"],
        ["Humanitarian cash and vouchers", "US$7.3B in 2025"],
        ["Aid already on-chain: WFP Building Blocks", "US$555M across 25M transactions"],
    ], [0.4, 0.6], size=14, row_height=Inches(0.5))

    # 4. What the officer sees
    s = d.slide(
        "3 · What Amanah does",
        "An officer asks in plain words, confirms one card, and the vault pays within the mandate",
        "docs/demo-transcript.md, steps 1 to 4; scenario 6 on Base Sepolia, {tx6}",
        "Here is a monthly aid run. The officer types: pay October's aid for Ipoh. There are two October batches, "
        "so the agent asks one question. The officer picks zakat monthly aid. The plan card shows 12 households "
        "and 1,810 USDC, built from the books. The officer says: leave out household 12, they moved. The agent "
        "re-plans the same batch: 11 households, 1,650 USDC. On confirm, the planner signs each payout, the auditor "
        "and the sentinel check it in parallel, and the vault decides the tier. All 11 are routine, so they are "
        "paid with no human in the loop, and every step goes to a hash-chained log.")
    steps = [("Understand", "one clarifying question"), ("Plan", "one card, amounts from the books"),
             ("Correct", "\"leave out household 12\""), ("Sign", "an EIP-712 intent"),
             ("Check", "auditor and sentinel, in parallel"), ("Submit", "the vault picks the tier"),
             ("Record", "a hash-chained log")]
    bw, gap = Inches(1.5), Inches(0.215)
    for i, (title, body) in enumerate(steps):
        x = MARGIN + i * (bw + gap)
        boxed(s, x, TOP, bw, Inches(1.3), [(title, 13, GREEN, True), (body, 12, INK, False)],
              align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            arrow(s, x + bw + Inches(0.02), TOP + Inches(0.52), width=Inches(0.18), height=Inches(0.24))
    convo = [("Officer", "Pay October's aid for Ipoh."),
             ("Amanah", "October has two batches for Ipoh: zakat monthly aid and flood relief. Which one?"),
             ("Officer", "Zakat monthly aid."),
             ("Amanah", "Plan: 12 households, 1,810 USDC. Tier preview: routine."),
             ("Officer", "Leave out household 12, they moved."),
             ("Amanah", "Same batch: 11 households, 1,650 USDC. Confirm?")]
    y = TOP + Inches(1.6)
    for who, said in convo:
        officer = who == "Officer"
        x = MARGIN + (Inches(0) if officer else Inches(1.2))
        boxed(s, x, y, Inches(10.9), Inches(0.42), [(f"{who}: {said}", 13, INK, officer)],
              fill_color=WHITE if officer else PALE, line=LINE if officer else None)
        y += Inches(0.5)

    # 5. Contribution 1
    s = d.slide(
        "4 · Contribution 1 of 5: the independence quorum",
        "Three signers built to fail differently must agree, and the contract checks they differ",
        "AmanahVault's signer registry and lineage check; scenario 11 on Base Sepolia, {tx11}",
        "A payout needs three signatures. The planner is Claude, by Anthropic, and it talks to staff. The auditor "
        "is Qwen, by Alibaba, a different model lineage. It never sees the conversation: it judges a case file "
        "of typed fields rebuilt from the books. The sentinel is not a language model at all. It applies eight "
        "rules to numbers, so no injected text can reach it. The signers differ in lineage, in what they see and "
        "in how they reason. The contract reads each key's lineage from an on-chain registry, and a quorum signed "
        "by two Claude keys reverts. Attacker text reaches at most one signer in raw form.")
    table(s, MARGIN, TOP, Inches(8.4), [
        ["Signer", "Lineage", "Sees", "Never sees"],
        ["Planner", "Claude, by Anthropic", "the chat; the books without personal data", "identity numbers, raw documents"],
        ["Auditor", "Qwen, by Alibaba", "the payout; a typed case file from the books", "the chat, notes, documents"],
        ["Sentinel", "8 rules, no model", "numbers only: the books, the batch, the vault's events", "any text"],
    ], [0.15, 0.22, 0.36, 0.27], size=14, row_height=Inches(1.05))
    boxed(s, MARGIN + Inches(8.75), TOP, Inches(3.35), Inches(4.2),
          [("Checked by the contract", 16, GREEN, True),
           ("Each key's lineage is in an on-chain registry that only two trustees can change.", 14, INK, False),
           ("A quorum signed by two Claude keys reverts with SameLineage: scenario 11.", 14, INK, False)],
          anchor=MSO_ANCHOR.TOP)

    # 6. Contributions 2 and 4
    s = d.slide(
        "5 · Contributions 2 and 4: tiers on-chain, bounded loss",
        "The vault, not the agent, decides when a human signs, and caps what machines can move",
        "tierOf() in contracts/src/AmanahVault.sol; forge test --match-contract Invariant; docs/contract-tests.md",
        "The contract decides how many signers a payout needs, from facts it can check itself: the amount against "
        "the programme's cap, the share of today's budget already spent, the household's spend this period, and "
        "the security posture. Routine aid "
        "flows with three machine signatures. Anything above the cap or past 80% of the day's budget needs a "
        "trustee too. An agent cannot argue its way into a lower tier, because the tier is never its decision. "
        "And the worst case is a number. With every machine signer compromised at once, money only reaches "
        "beneficiaries two trustees approved, at most 80% of each daily budget moves per day, and each household "
        "stays within its 30-day cap. Eight invariants hold over 16,384 random calls each.")
    table(s, MARGIN, TOP, Inches(7.6), [
        ["Tier", "When", "Who signs"],
        ["1 Routine", "at or under the cap, today's spend at or under 80% of the budget, posture normal",
         "planner, auditor, sentinel"],
        ["2 Elevated", "above the cap, past the 80% line, or posture raised", "the three, plus a trustee"],
        ["3 Override", "citing an anchored refusal for the same payout", "planner and two trustees"],
    ], [0.2, 0.5, 0.3], size=14, row_height=Inches(1.05))
    boxed(s, MARGIN + Inches(7.95), TOP, Inches(4.15), Inches(4.2),
          [("If every machine signer is compromised", 16, GREEN, True),
           ("Money reaches only beneficiaries two trustees approved for that programme.", 14, INK, False),
           ("At most 80% of each daily budget moves without a human, per day.", 14, INK, False),
           ("Each household stays within its 30-day cap.", 14, INK, False),
           ("8 invariants hold over 16,384 random calls each.", 14, INK, True)], anchor=MSO_ANCHOR.TOP)

    # 7. Contribution 3
    s = d.slide(
        "6 · Contribution 3 of 5: the immune loop",
        "When attacks start to work, the wallet gets stricter on its own, and only two trustees can loosen it",
        "redteam/immune.py and backend/app/agentwallet/posture.py; scenario 14 on Base Sepolia, {tx14}",
        "Models change without notice and new jailbreaks keep appearing, so a control tested once decays. Every "
        "day a red-team run fires the reviewed attack gallery at a sandbox copy of the quorum. It measures how "
        "often the planner and the auditor fail together, and it anchors the report's hash on-chain. Above a "
        "threshold fixed before the run, it raises the posture, and every payout then needs a trustee. The red "
        "team can only raise the posture. Lowering it takes two trustee signatures. So a compromised red-team "
        "run can cause delays, never theft. Our measured joint failure was 0.0%, which raises nothing, so the "
        "demo's scenario 14 uses a summary labelled as made for the demo.")
    cycle = [("Attack", "a sandbox copy, every day"), ("Measure", "joint failure of planner and auditor"),
             ("Anchor", "the report's hash, on-chain"), ("Raise", "posture 1 above 10%, posture 2 above 25%"),
             ("Hold", "every payout needs a trustee"), ("Lower", "only with two trustee signatures")]
    bw, gap = Inches(1.8), Inches(0.26)
    for i, (title, body) in enumerate(cycle):
        x = MARGIN + i * (bw + gap)
        color = AMBER if title in ("Raise", "Hold") else GREEN
        boxed(s, x, TOP + Inches(0.2), bw, Inches(1.9), [(title, 17, color, True), (body, 14, INK, False)],
              align=PP_ALIGN.CENTER)
        if i < len(cycle) - 1:
            arrow(s, x + bw + Inches(0.02), TOP + Inches(1.0), width=Inches(0.22), height=Inches(0.28))
    boxed(s, MARGIN, TOP + Inches(2.55), W - 2 * MARGIN, Inches(1.7),
          [("A one-way ratchet", 16, GREEN, True),
           ("The posture oracle's key can only anchor reports and raise the posture. Lowering it needs two trustee "
            "signatures, which the vault checks itself.", 15, INK, False),
           ("The thresholds were fixed in redteam/protocol.md before the full run.", 15, INK, False)],
          anchor=MSO_ANCHOR.TOP)

    # 8. The four rubric attacks
    s = d.slide(
        "7 · Security: the four attacks in the rubric",
        "Each attack in the rubric meets at least three layers, and the last layer is code",
        "docs/security-self-assessment.md, section 3; scenarios 10 {tx10} and 12 {tx12}; the live planner refused 8 and 9",
        "The rubric names four attacks. An induced transfer, such as an urgent message from the chairman, meets "
        "an auditor that finds no such beneficiary, a sentinel that sees an unknown wallet, and a vault that "
        "reverts anything off the allowlist. A prompt-injected letter reaches the planner only as typed fields, "
        "the amount still comes from the books, and the sentinel's rules refuse the second payment it asks for. "
        "On Base Sepolia the live planner refused both of those messages itself, so the red team shows the "
        "checkers instead: every checker rejected the payouts to unlisted wallets that the planner proposed. "
        "Data extraction finds nothing to leak, because no model's view holds personal data. Privilege escalation "
        "fails because no agent key holds a trustee role, and every registry change needs two trustees.")
    table(s, MARGIN, TOP, W - 2 * MARGIN, [
        ["Attack", "The layers it meets", "Shown by"],
        ["Induced transfer", "the auditor finds no such beneficiary; the sentinel sees an unknown wallet; the vault "
         "reverts any wallet off the allowlist", "scenario 10; the red team's induced-transfer and urgency cases"],
        ["Prompt injection", "a quarantined reader returns typed fields only; the amount comes from the books; the "
         "sentinel's frequency and structuring rules refuse", "scenario 9"],
        ["Data extraction", "no personal data in any model's view; on-chain, a beneficiary is a wallet and a keyed "
         "hash", "the red team's extraction cases"],
        ["Privilege escalation", "no agent key holds a trustee role; every registry, mandate and allowlist change "
         "needs two trustee signatures", "scenario 12"],
    ], [0.18, 0.6, 0.22], size=14, row_height=Inches(0.85))

    # 9. Contribution 5
    s = d.slide(
        "8 · Contribution 5 of 5: measured independence",
        "We measured how often the signers fail together, and we report null results as they are",
        "redteam/protocol.md, committed at {protocol_commit} before the run; redteam/results/final/report.md; "
        "{n_attacks} attacks and {n_legit} legitimate payouts",
        "Independence is the whole argument, so we measured it rather than assumed it. We fixed the protocol and "
        "committed it before the full run. Each attack ran against five configurations, from the planner alone up "
        "to the full quorum and the vault. Against the planner alone, {S1_short} of the attacks succeeded. With "
        "the Qwen auditor, {S3_short}. With the sentinel too, {S4_short}. Through the vault, {S5_short}. The planner and "
        "the Qwen auditor failed together on {joint} of transfer attacks. The paired comparisons use exact McNemar "
        "tests with Holm's correction, and every rate has a 95% Wilson interval. One caution: the single attack "
        "that passed the Claude auditor but not the Qwen auditor was stopped by a Qwen misreading of its own "
        "rules, the same misreading that escalated two sound payouts. So that difference rests on one case, with "
        "a McNemar p of 1.0, and we do not claim Qwen judges better. We did not tune the prompt to those cases.")
    table(s, MARGIN, TOP, Inches(8.6), [
        ["Arm", "What must sign before money moves", "Attack success, all 44", "False refusals"],
        ["S1", "the planner alone", d.s("{S1}"), d.s("{FR1}")],
        ["S2", "the planner and a Claude auditor", d.s("{S2}"), d.s("{FR2}")],
        ["S3", "the planner and the Qwen auditor", d.s("{S3}"), d.s("{FR3}")],
        ["S4", "S3 and the sentinel", d.s("{S4}"), d.s("{FR4}")],
        ["S5", "S4 and the vault, on Anvil", d.s("{S5}"), d.s("{FR5}")],
    ], [0.08, 0.37, 0.33, 0.22], size=13, row_height=Inches(0.7))
    boxed(s, MARGIN + Inches(8.95), TOP, Inches(3.15), Inches(4.2),
          [("Joint failure", 16, GREEN, True),
           (d.s("The planner and the Qwen auditor were fooled together on {joint} of transfer attacks."), 14, INK,
            False),
           ("Above 10% raises the posture, by the protocol.", 14, INK, False),
           ("S2 against S3 is one case, p = 1.0: not a claim that Qwen judges better.", 13, MUTED, False)],
          anchor=MSO_ANCHOR.TOP)

    # 10. Task completion
    s = d.slide(
        "9 · Task completion on Base Sepolia",
        "All 16 scripted scenarios ran on Base Sepolia, and any executed payout re-checks from the chain alone",
        "evidence/transactions.csv, evidence/demo-summary.json; python scripts/verify.py --csv evidence/transactions.csv",
        "Task completion is forty percent of the score, so every scenario ran on Base Sepolia in USDC. The vault "
        "is at {vault}. Scenarios one to five set it up: two trustees register the signers, set three mandates and "
        "add 45 beneficiaries, and a donor deposits. Then come the payouts, the attacks and the controls. In all, "
        "{tx_count} transactions are tagged with their scenario, and {scenarios_passed} of 16 scenarios behaved as "
        "expected. Our verify script re-checks any executed payout from the chain alone: signer roles, distinct "
        "lineages, the tier, the caps, the allowlist and the trustee signature. It passed on {verified} executed "
        "payouts, and it fails on a hand-edited one.")
    names = ["Deploy the vault", "Register signers", "Set 3 mandates", "Add 45 beneficiaries", "Donor deposit",
             "Routine batch of 11", "Grant with a trustee", "Induced transfer", "Injected letter",
             "Forced, off the allowlist", "Same lineage", "Planner governs", "Structuring", "Immune loop", "Override",
             "Guardian pause"]
    outcome = ["deployed", "registered", "set", "added", "received", "paid", "paid", "refused, anchored",
               "refused, anchored", "reverted", "reverted", "reverted", "capped", "posture raised", "paid", "reverted"]
    cw, ch, gx, gy = Inches(1.9), Inches(0.92), Inches(0.12), Inches(0.12)
    for i, name in enumerate(names):
        col, row = i % 4, i // 4
        x, y = MARGIN + col * (cw + gx), TOP + row * (ch + gy)
        status = d.f[f"ok{i + 1}"]
        done = "refused by the planner" if d.f[f"tx{i + 1}"] == "no transaction" else outcome[i]
        boxed(s, x, y, cw, ch, [(f"{i + 1}  {name}", 12, INK, True), (f"{done} · {status}", 11, MUTED, False)])
    boxed(s, MARGIN + 4 * (cw + gx) + Inches(0.2), TOP, Inches(3.65), 4 * ch + 3 * gy,
          [("On Base Sepolia", 16, GREEN, True),
           (d.s("Vault {vault}"), 14, INK, False),
           (d.s("{tx_count} transactions, each tagged with its scenario"), 14, INK, False),
           (d.s("verify.py passes on {verified} executed payouts"), 14, INK, False),
           ("The action log is hash-chained, and its head is anchored on-chain.", 14, INK, False)],
          anchor=MSO_ANCHOR.TOP)

    # 11. Path to deployment
    s = d.slide(
        "10 · From demo to pilot",
        "The same wallet serves any fund whose rules fit a mandate, and the path to a pilot is short",
        "README sections 7.6, 11 and 16; bash scripts/test_all.sh",
        "Amanah is fund-agnostic: a programme's rules live in its on-chain mandate, so the same wallet serves "
        "zakat, waqf, charity, humanitarian cash or a government transfer. The demo runs three funds side by side. "
        "The build has {tests_total} automated tests across six suites. A pilot needs one programme at a "
        "Malaysian zakat institution, agent keys in a key management service, and trustee governance in a Safe "
        "multisig. After that come attested lineage, so the registry can prove which model signed, and a "
        "permissioned chain for institutions that need one. The design maps to China's 2026 agent-payment "
        "convention: authorisation boundaries, limits, Know Your Agent, security testing and a human fallback.")
    columns = [("Built and tested", ["3 funds, 3 programmes, 45 beneficiaries", "16 scripted scenarios",
                                     d.s("{tests_total} automated tests: {tests_split}")]),
               ("For a pilot", ["one programme at a Malaysian zakat institution, planned outreach",
                                "agent keys in a key management service", "trustee governance in a Safe multisig"]),
               ("Next", ["attested lineage, so the registry proves which model signed",
                         "FISCO BCOS beyond one node, after 5 of 5 checks on a local chain",
                         "shared keyed hashes, so institutions avoid paying one need twice"])]
    for i, (title, items) in enumerate(columns):
        x = MARGIN + i * Inches(4.1)
        boxed(s, x, TOP, Inches(3.85), Inches(2.85), [(title, 17, GREEN, True)] + [(t, 14, INK, False) for t in items],
              anchor=MSO_ANCHOR.TOP)
    boxed(s, MARGIN, TOP + Inches(3.1), W - 2 * MARGIN, Inches(0.95),
          [("Compliance: the 2026 agent-payment convention asks for authorisation boundaries, transaction limits, "
            "Know Your Agent, security testing and human fallback. Each maps to a control here (README 7.6).",
            14, INK, False)], fill_color=WHITE, line=LINE)
    return d


def fills_left(deck_path: Path) -> list[str]:
    import zipfile

    found = []
    with zipfile.ZipFile(deck_path) as z:
        for name in sorted(z.namelist()):
            if re.fullmatch(r"ppt/(slides|notesSlides)/\w+\.xml", name):
                found += re.findall(r"\[\[fill: ([^\]]+)\]\]", z.read(name).decode())
    return sorted(set(found))


def to_pdf(pptx: Path) -> Path:
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office is None:
        raise SystemExit("LibreOffice is needed for --pdf: install it, or export the PDF from PowerPoint")
    subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)],
                   check=True, capture_output=True, timeout=300)
    return pptx.with_suffix(".pdf")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python scripts/build_deck.py", description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", type=Path, default=OUT)
    parser.add_argument("--pdf", action="store_true", help="also export a PDF with LibreOffice")
    parser.add_argument("--team")
    parser.add_argument("--members")
    args = parser.parse_args(argv)
    deck = build(facts(args.team, args.members))
    deck.save(args.out)
    print(f"{deck.count} slides written to {args.out}")
    if args.pdf:
        print(f"PDF written to {to_pdf(args.out)}")
    left = fills_left(args.out)
    if left:
        print(f"{len(left)} values still to fill, from the runs or from Hafiz: " + "; ".join(left))
    return 0


if __name__ == "__main__":
    sys.exit(main())
